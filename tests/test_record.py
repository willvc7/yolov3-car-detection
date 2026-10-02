# -*- coding: utf-8 -*-
"""record schema parity + train-time column behavior (tmp workbooks, no Drive)."""

import json
from pathlib import Path

from car_tools import record as R

NODRIVE = "D:\\__no_drive__"


def _wb(exp, tmp_path, **kw):
    runs = Path(tmp_path) / "runs"
    return R.update_workbook(exp, runs_root=str(runs),
                             cloud_xlsx=str(runs / "cloud.xlsx"),
                             local_xlsx=str(runs / "rec.xlsx"),
                             drive_marker=NODRIVE, **kw)


def test_total_h_g_parity():
    assert len(R.TOTAL_H) == len(R.TOTAL_G) == 37, (len(R.TOTAL_H), len(R.TOTAL_G))
    assert R.TOTAL_H.index("training_time(h)") == R.TOTAL_H.index("mAP50-95") + 1
    # group spans must tile exactly: recompute run-lengths of TOTAL_G
    runs, start = [], 0
    for i in range(1, len(R.TOTAL_G) + 1):
        if i == len(R.TOTAL_G) or R.TOTAL_G[i] != R.TOTAL_G[start]:
            runs.append((R.TOTAL_G[start], start, i - 1))
            start = i
    assert sum(e - s + 1 for _, s, e in runs) == len(R.TOTAL_H)
    by_group = {g: [R.TOTAL_H[c] for c in range(s, e + 1)] for g, s, e in runs}
    assert by_group["訓練設定"] == ["epochs", "batch", "imgsz"]
    assert by_group["成績 metrics"] == ["P", "R", "mAP50", "mAP50-95", "training_time(h)"]
    assert by_group["驗證推論條件"] == ["驗證iou", "推論conf"]
    assert len(R.CLS_H) == len(R.CLS_G) == 8


def test_train_time_rounding(tmp_path):
    runs = tmp_path / "runs"
    (runs / "train" / "exp_T").mkdir(parents=True)
    (runs / "train" / "exp_T" / "train_time.json").write_text(
        json.dumps({"seconds": 1274.4}), encoding="utf-8")
    row = _wb("exp_T", tmp_path)
    assert row["training_time(h)"] == 0.354


def test_train_time_missing_blank(tmp_path):
    row = _wb("exp_T", tmp_path)
    assert row.get("training_time(h)") is None


def test_train_time_manual_preserved(tmp_path):
    runs = tmp_path / "runs"
    (runs / "train" / "exp_T").mkdir(parents=True)
    (runs / "train" / "exp_T" / "train_time.json").write_text(
        json.dumps({"seconds": 7200}), encoding="utf-8")
    _wb("exp_T", tmp_path)  # writes 2.0
    (runs / "train" / "exp_T" / "train_time.json").unlink()  # json gone
    row = _wb("exp_T", tmp_path)  # must keep 2.0, not blank it
    assert row["training_time(h)"] == 2.0


def test_train_time_new_value_overwrites(tmp_path):
    runs = tmp_path / "runs"
    (runs / "train" / "exp_T").mkdir(parents=True)
    (runs / "train" / "exp_T" / "train_time.json").write_text(
        json.dumps({"seconds": 7200}), encoding="utf-8")
    _wb("exp_T", tmp_path)
    (runs / "train" / "exp_T" / "train_time.json").write_text(
        json.dumps({"seconds": 3600}), encoding="utf-8")
    row = _wb("exp_T", tmp_path)
    assert row["training_time(h)"] == 1.0


def _mk_val_withOpt(tmp_path, val_body, opt_data="data/car.yaml", names=None):
    runs = tmp_path / "runs"
    (runs / "val" / "exp_T").mkdir(parents=True)
    (runs / "val" / "exp_T" / "val.txt").write_text(val_body, encoding="utf-8")
    (runs / "train" / "exp_T").mkdir(parents=True)
    (runs / "train" / "exp_T" / "opt.yaml").write_text(
        "data: %s\n" % opt_data, encoding="utf-8")
    if names is not None:
        (tmp_path / "data").mkdir(exist_ok=True)
        import yaml as _y
        (tmp_path / "data" / "car.yaml").write_text(
            _y.dump({"names": names}, allow_unicode=True), encoding="utf-8")


def _read_cls(tmp_path):
    import openpyxl as _ox
    wb = _ox.load_workbook(tmp_path / "runs" / "rec.xlsx")
    if "類別明細" not in wb.sheetnames:
        return None
    ws = wb["類別明細"]
    return [[ws.cell(r, c).value for c in range(1, 9)]
            for r in range(3, ws.max_row + 1)]


def test_class_rows_multiclass_parsed(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _mk_val_withOpt(tmp_path,
        "                   all         71        107      0.98      0.95      0.98      0.57\n"
        "                 car         60         90      0.99      0.96      0.99      0.60\n"
        "                 bus         11         17      0.90      0.80      0.85      0.40\n")
    _wb("exp_T", tmp_path)
    rows = _read_cls(tmp_path)
    assert rows is not None and len(rows) == 2
    assert rows[0][:3] == ["exp_T", "car", 90]
    assert rows[0][3:7] == [0.99, 0.96, 0.99, 0.6]


def test_class_rows_synthesized_single_class(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _mk_val_withOpt(tmp_path,
        "                   all         71        107    0.00465      0.925      0.396      0.142\n",
        names=["car"])
    _wb("exp_T", tmp_path)
    rows = _read_cls(tmp_path)
    assert rows is not None and len(rows) == 1
    assert rows[0][:3] == ["exp_T", "car", 107]
    assert rows[0][3:7] == [0.00465, 0.925, 0.396, 0.142]


def test_class_rows_not_synthesized_multi_class(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _mk_val_withOpt(tmp_path,
        "                   all         71        107      0.98      0.95      0.98      0.57\n",
        names=["car", "bus"])
    _wb("exp_T", tmp_path)
    assert _read_cls(tmp_path) is None


def test_class_rows_fallback_all_when_names_missing(tmp_path, monkeypatch):
    # 真實 val.txt 只有 all 列、且 data yaml 找不到時：保留 all 聚合值而非靜默丟失
    monkeypatch.chdir(tmp_path)
    runs = tmp_path / "runs"
    (runs / "val" / "exp_T").mkdir(parents=True)
    (runs / "val" / "exp_T" / "val.txt").write_text(
        "                   all         71        107    0.00465      0.925      0.396      0.142\n",
        encoding="utf-8")
    (runs / "train" / "exp_T").mkdir(parents=True)  # 無 opt.yaml -> names 無法解析
    _wb("exp_T", tmp_path)
    rows = _read_cls(tmp_path)
    assert rows is not None and len(rows) == 1
    assert rows[0][:3] == ["exp_T", "all", 107]
    assert rows[0][3:7] == [0.00465, 0.925, 0.396, 0.142]
