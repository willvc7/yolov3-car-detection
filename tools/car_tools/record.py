# -*- coding: utf-8 -*-
"""Experiment + per-class metrics workbook maintenance (idempotent).

Ported from YOLOv3_try2.ipynb Cell47. Statements unchanged;
paths/conditions became parameters (defaults preserve Cell47 behavior).
"""

import math
import re
from pathlib import Path

import pandas as pd
import yaml

try:
    import openpyxl
except ImportError:  # pragma: no cover
    raise SystemExit("缺少 openpyxl，請先安裝（pip install openpyxl）")

TOTAL_H = ["實驗名稱", "日期", "資料集", "epochs", "epochs_done", "batch", "imgsz",
           "lr0", "lrf", "momentum", "weight_decay", "warmup_epochs",
           "box", "cls", "obj", "iou_t", "anchor_t", "mosaic", "mixup", "fl_gamma", "scale",
           "初始weights", "cfg", "hyp檔",
           "驗證iou", "驗證用權重", "推論conf",
           "train_box_loss", "train_obj_loss", "train_cls_loss",
           "val_box_loss", "val_obj_loss", "val_cls_loss",
           "P", "R", "mAP50", "mAP50-95",
           "best.pt雲端路徑", "備份資料夾", "備註"]
TOTAL_G = ["基本資訊", "基本資訊", "基本資訊",
           "訓練設定", "訓練設定", "訓練設定", "訓練設定",
           "學習率排程", "學習率排程", "學習率排程", "學習率排程", "學習率排程",
           "損失與錨點", "損失與錨點", "損失與錨點", "損失與錨點", "損失與錨點",
           "資料增強", "資料增強", "資料增強", "資料增強",
           "模型與資料", "模型與資料", "模型與資料",
           "驗證推論條件", "驗證推論條件", "驗證推論條件",
           "損失終值", "損失終值", "損失終值", "損失終值", "損失終值", "損失終值",
           "成績 metrics", "成績 metrics", "成績 metrics", "成績 metrics",
           "檔案位置", "檔案位置", "備註"]
CLS_G = ["基本資訊", "類別資訊", "類別資訊",
         "成績 metrics", "成績 metrics", "成績 metrics", "成績 metrics", "解讀"]
CLS_H = ["實驗名稱", "類別", "Instances(數量)",
         "P", "R", "mAP50", "mAP50-95", "解讀(初學者看這欄)"]
HYP_KEYS = ["lr0", "lrf", "momentum", "weight_decay", "warmup_epochs",
            "box", "cls", "obj", "iou_t", "anchor_t", "mosaic", "mixup", "fl_gamma", "scale"]
MET_KEYS = [("metrics/precision", "P"), ("metrics/recall", "R"),
            ("metrics/mAP_0.5", "mAP50"), ("metrics/mAP_0.5:0.95", "mAP50-95")]
LOSS_KEYS = [("train/box_loss", "train_box_loss"), ("train/obj_loss", "train_obj_loss"),
             ("train/cls_loss", "train_cls_loss"), ("val/box_loss", "val_box_loss"),
             ("val/obj_loss", "val_obj_loss"), ("val/cls_loss", "val_cls_loss")]

HEADER_ROW = 2  # 第1列：分類色塊；第2列：正式欄位名稱

DATA_DISPLAY = {"data/car.yaml": "Car-Object-Detection(1類)"}


def _merge_runs(ws, groups):
    """合併第1列相同群組的儲存格（安全版，重複執行不報錯）"""
    start = 0
    for i in range(1, len(groups) + 1):
        if i == len(groups) or groups[i] != groups[start]:
            if i - 1 > start:
                try:
                    ws.merge_cells(start_row=1, start_column=start + 1,
                                   end_row=1, end_column=i)
                except Exception:
                    pass  # 已合併則跳過
            start = i


def ensure(ws_name, headers, groups, wb):
    """確保工作表存在且有正確的雙行表頭，可重複執行。"""
    if ws_name not in wb.sheetnames:
        ws = wb.create_sheet(ws_name)
    else:
        ws = wb[ws_name]
        # 表頭全欄比對（只比第一欄會導致錯位 append）
        if ws.max_row >= 2 and [ws.cell(2, c).value for c in range(1, len(headers) + 1)] == list(headers):
            return ws
        _has = ws.max_row > HEADER_ROW and any(
            ws.cell(r, 1).value is not None for r in range(HEADER_ROW + 1, ws.max_row + 1))
        if _has:
            raise ValueError(f"工作表「{ws_name}」表頭與程式預期不符，為避免錯位寫入請先備份並對齊表頭")
    # 建立雙行表頭（群組列 + 欄位列）
    for c, g in enumerate(groups, 1):
        ws.cell(1, c).value = g
    _merge_runs(ws, groups)
    for c, h in enumerate(headers, 1):
        ws.cell(2, c).value = h
    return ws


def update_workbook(exp, runs_root="runs",
                    cloud_xlsx="/content/drive/MyDrive/YOLO_Experiments/數據紀錄表.xlsx",
                    local_xlsx="數據紀錄表.xlsx",
                    drive_marker="/content/drive/MyDrive",
                    val_iou=0.65, val_w="best.pt", det_conf=0.25, note=None):
    """收集 runs/<exp> 產物並寫入（可重複跑不重複寫）。note 寫入備註欄。"""
    assert exp, "缺少 EXP_NAME"
    print("[record v6.0] EXP =", exp)
    runs_root = Path(runs_root)
    results = runs_root / "train" / exp / "results.csv"
    opt = runs_root / "train" / exp / "opt.yaml"
    hyp_used = runs_root / "train" / exp / "hyp.yaml"
    valtxt = runs_root / "val" / exp / "val.txt"
    hsrc = runs_root / "train" / exp / "hyp_source.txt"

    # ── 收集資料 ──────────────────────────────────────────────
    row = {"實驗名稱": exp,
           "備份資料夾": f"YOLO_Experiments/{exp}",
           "best.pt雲端路徑": f"YOLO_Experiments/{exp}/train/{exp}/weights/best.pt",
           "驗證iou": val_iou, "驗證用權重": val_w, "推論conf": det_conf}
    if note:
        row["備註"] = note

    if results.exists():
        df = pd.read_csv(results)
        df.columns = [c.strip() for c in df.columns]
        last = df.iloc[-1].to_dict()
        for k, nk in MET_KEYS + LOSS_KEYS:
            if k in last:
                try:
                    _v = round(float(last[k]), 4)
                    if isinstance(_v, float) and math.isnan(_v):
                        continue  # 中斷訓練的殘缺列不寫入，避免 NaN 污染表格
                    row[nk] = _v
                except (TypeError, ValueError):
                    pass
        row["epochs_done"] = len(df)
        row.setdefault("epochs", len(df))
    else:
        print(f"找不到 {results}，先確認訓練已完成")

    if opt.exists():
        with open(opt, encoding="utf-8") as _fh:
            d = yaml.safe_load(_fh) or {}
        for k, nk in {"batch_size": "batch", "imgsz": "imgsz", "cfg": "cfg", "data": "資料集"}.items():
            if k in d and d[k] not in (None, ""):
                row[nk] = d[k]
        if row.get("資料集") in DATA_DISPLAY:
            row["資料集"] = DATA_DISPLAY[row["資料集"]]
        if "epochs" in d:
            try:
                row["epochs"] = int(d["epochs"])
            except (TypeError, ValueError):
                pass
        w = str(d.get("weights", ""))
        row["初始weights"] = w if w else "從零訓練(scratch)"
        h = d.get("hyp", "")
        if isinstance(h, dict):
            for k in HYP_KEYS:
                if k in h:
                    row[k] = h[k]
            row["hyp檔"] = "(opt.yaml內嵌dict)"
        elif h:
            row["hyp檔"] = Path(str(h)).name
        if hsrc.exists():
            _hs = hsrc.read_text(encoding="utf-8").strip().splitlines()[0]
            if isinstance(h, str) and h and Path(str(h)).name != Path(_hs).name:
                print(f"hyp 來源不一致：opt.yaml 記 {h}，hyp_source.txt 記 {_hs}（以 hyp_source.txt 為準，請確認訓練參數）")
            row["hyp檔"] = Path(_hs).name
        m0 = re.match(r"exp_(\d{4})(\d{2})(\d{2})_", exp)
        row["日期"] = f"{m0.group(1)}-{m0.group(2)}-{m0.group(3)}" if m0 else ""

    if hyp_used.exists():
        with open(hyp_used, encoding="utf-8") as _fh2:
            h = yaml.safe_load(_fh2) or {}
        for k in HYP_KEYS:
            if k in h and k not in row:
                row[k] = h[k]

    print("待寫入總表:", row)

    # ── 解析 val.txt 類別明細 ─────────────────────────────────
    cls_rows = []
    if valtxt.exists():
        t = valtxt.read_text(encoding="utf-8", errors="ignore").replace("\r", "\n")
        t = re.sub(r"\x1b\[[0-9;]*m", "", t)
        for line in t.splitlines():
            p = line.strip().split()
            if len(p) >= 7 and p[1].isdigit() and p[2].isdigit() and p[0] != "all":
                try:
                    vals = [float(x) for x in p[3:7]]
                except ValueError:
                    continue
                cls_rows.append({"實驗名稱": exp, "類別": p[0],
                                 "Instances(數量)": int(p[2]),
                                 "P": vals[0], "R": vals[1],
                                 "mAP50": vals[2], "mAP50-95": vals[3],
                                 "解讀(初學者看這欄)": ""})
        print(f"解析 val.txt：{len(cls_rows)} 個類別")
    else:
        print(f"找不到 {valtxt}，跳過類別明細")

    # ── 寫入 Excel ────────────────────────────────────────────
    if Path(drive_marker).exists():
        target = Path(cloud_xlsx)
    else:
        target = Path(local_xlsx)
        print("Drive 未掛載，紀錄表寫入本地（之後需手動合併，避免與雲端表分叉）")
    wb = openpyxl.load_workbook(target) if target.exists() else openpyxl.Workbook()

    ws = ensure("實驗總表", TOTAL_H, TOTAL_G, wb)
    _vals = [row.get(h) for h in TOTAL_H]
    _row_idx = next(
        (r for r in range(HEADER_ROW + 1, ws.max_row + 1) if str(ws.cell(r, 1).value) == exp), None)
    if _row_idx is not None:
        for _c, _v in enumerate(_vals, 1):
            ws.cell(_row_idx, _c).value = _v
        print(f"{exp} 已存在，已原地更新第 {_row_idx} 列（重跑會覆寫，不會重複）")
    else:
        ws.append(_vals)
        print(f"已追加到實驗總表，共 {ws.max_row - HEADER_ROW} 筆")

    if cls_rows:
        ws2 = ensure("類別明細", CLS_H, CLS_G, wb)
        _idx = {}
        for r in range(HEADER_ROW + 1, ws2.max_row + 1):
            if ws2.cell(r, 1).value is not None:
                _idx[(str(ws2.cell(r, 1).value), str(ws2.cell(r, 2).value))] = r
        n_new, n_upd = 0, 0
        for cr in cls_rows:
            _vals2 = [cr.get(h) for h in CLS_H]
            _r = _idx.get((exp, cr["類別"]))
            if _r is None:
                ws2.append(_vals2)
                n_new += 1
            else:
                for _c, _v in enumerate(_vals2, 1):
                    ws2.cell(_r, _c).value = _v
                n_upd += 1
        print(f"類別明細新增 {n_new} 列、更新 {n_upd} 列")

    # 清除預設空白 Sheet
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"].cell(1, 1).value is None:
        del wb["Sheet"]

    Path(target).parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    print(f"已儲存 {target}")
    return row