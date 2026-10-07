# -*- coding: utf-8 -*-
"""val-iou x detect-conf sweep with caching + Excel logging + cloud backup.

Ported from YOLOv3_try2.ipynb Cell40. Statements unchanged;
notebook globals (exp/backup_dir/W/SRC/paths) became parameters.
"""

import json
import re
import subprocess
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    raise SystemExit("缺少 openpyxl，請先安裝（pip install openpyxl）")

from .exp import CLOUD_XLSX, DRIVE_MARKER, LOCAL_XLSX

SWP_G = ["基本資訊", "掃參條件", "掃參條件",
         "成績 metrics", "成績 metrics", "成績 metrics", "成績 metrics",
         "結果備註", "結果備註"]
SWP_H = ["實驗名稱", "掃參類型(val-iou/detect-conf)", "參數值",
         "P", "R", "mAP50", "mAP50-95", "檢出摘要", "備註"]
HEADER_ROW = 2

_ansi = re.compile(r"\x1b\[[0-9;]*m")


def run(cmd, timeout=3600):
    """執行子程序，回傳合併後的純文字輸出（已去除 ANSI 色碼）"""
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired as e:
        raise RuntimeError(f"指令逾時（>{timeout}s）：{' '.join(cmd)}") from e
    t = _ansi.sub("", (p.stdout or "") + "\n" + (p.stderr or "")).replace("\r", "\n")
    if p.returncode != 0:
        raise RuntimeError(
            f"指令失敗 (returncode={p.returncode})：{' '.join(cmd)}\n"
            f"--- 輸出尾部 ---\n{t[-3000:]}"
        )
    return t


def all_row(t):
    """從 val.py 輸出中解析 'all' 那列的 P/R/mAP50/mAP50-95"""
    m = re.search(r"^\s*all\s+\d+\s+\d+\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)", t, re.M)
    return [round(float(x), 4) for x in m.groups()] if m else [None] * 4


def weights_key(p):
    """權重指紋（大小+mtime）：重訓後舊快取自動失效"""
    st = Path(p).stat()
    return f"{st.st_size}:{st.st_mtime_ns}"


def _meta_ok(mpath, kind, val, wk, weights):
    """快取有效 only if 參數+權重指紋都吻合（舊版無 meta.json 的快取視為失效，會重跑一次）"""
    try:
        m = json.loads(Path(mpath).read_text(encoding="utf-8"))
    except Exception:
        return False
    return (m.get("kind") == kind and str(m.get("param")) == str(val)
            and m.get("weights_key") == wk and m.get("weights") == weights)


def _meta_write(mpath, kind, val, wk, weights):
    Path(mpath).write_text(json.dumps(
        {"kind": kind, "param": str(val), "weights": weights, "weights_key": wk},
        ensure_ascii=False), encoding="utf-8")


def _require_metrics(t, label):
    vals = all_row(t)
    if None in vals:
        raise RuntimeError(
            f"{label} 的 val.py 輸出解析失敗（找不到 all 列），請檢查上方錯誤訊息\n"
            f"--- 輸出尾部 ---\n{t[-3000:]}")
    return vals


def _merge_runs(ws, groups):
    start = 0
    for i in range(1, len(groups) + 1):
        if i == len(groups) or groups[i] != groups[start]:
            if i - 1 > start:
                try:
                    ws.merge_cells(start_row=1, start_column=start + 1,
                                   end_row=1, end_column=i)
                except Exception:
                    pass
            start = i


def pick_sweep_image(val_images_dir):
    """取 val 首張當掃參代表圖。"""
    from glob import glob as _g
    cands = sorted(_g(str(Path(val_images_dir) / "*")))
    assert cands, f"{val_images_dir} 是空的，請先完成標註轉換"
    src = cands[0]
    print("[TRY] sweep SRC =", src)
    return src


def run_sweep(exp, weights, data="data/car.yaml", runs_root="runs",
              img=640, batch_size=16,
              iou_list=("0.5", "0.6", "0.65"), conf_list=("0.15", "0.25", "0.4"),
              val_images_dir="../datasets/car/images/val", src_image=None,
              cloud_xlsx=None, local_xlsx=None,
              backup_dir="", drive_marker=None):
    """掃 val-iou × detect-conf，備份輸出並寫入掃參紀錄。回傳 rows（供測試/除錯）。

    路徑固定於 exp 模組 (DRIVE_ROOT 固定): 呼叫方不傳即用預設值.
    """
    cloud_xlsx = cloud_xlsx or CLOUD_XLSX
    local_xlsx = local_xlsx or LOCAL_XLSX
    drive_marker = drive_marker or DRIVE_MARKER
    assert exp, "缺少 EXP_NAME"
    assert backup_dir, "缺少 BACKUP_DIR"
    print("[sweep v6.0] EXP =", exp)
    weights = str(weights)
    assert Path(weights).exists(), f"找不到權重 {weights}，請先完成訓練"
    runs_root = Path(runs_root)

    rows = []

    # ── val-iou 掃參 ──────────────────────────────────────────
    wk = weights_key(weights)  # 權重指紋：重訓後快取自動失效
    for iou in iou_list:
        iou = str(iou)
        name = f"{exp}_iou{iou.replace('.', '')}"
        vt = runs_root / "val" / name / "val.txt"
        mt = runs_root / "val" / name / "meta.json"
        t = None
        if vt.exists() and _meta_ok(mt, "val-iou", iou, wk, weights):
            t = _ansi.sub("", vt.read_text(encoding="utf-8", errors="ignore")).replace("\r", "\n")
            _vals = all_row(t)
            if _vals is None or None in _vals:
                print(f"iou={iou} 快取損毀（解析失敗），重新執行")
                t = None
            else:
                P, R, M50, M95 = _vals
                print(f"iou={iou} (沿用既有) P={P} R={R} mAP50={M50} mAP50-95={M95}")
        if t is None:
            t = run(["python", "val.py", "--img", str(img), "--batch-size", str(batch_size),
                     "--weights", weights, "--data", data,
                     "--iou", iou, "--project", str(runs_root / "val"), "--name", name, "--exist-ok"])
            (runs_root / "val" / name).mkdir(parents=True, exist_ok=True)
            vt.write_text(t, encoding="utf-8", errors="ignore")
            _meta_write(mt, "val-iou", iou, wk, weights)
            P, R, M50, M95 = _require_metrics(t, f"val-iou={iou}")
            print(f"iou={iou} P={P} R={R} mAP50={M50} mAP50-95={M95}")
        rows.append([exp, "val-iou", float(iou), P, R, M50, M95, "", f"權重 {weights}"])

    # ── detect-conf 掃參 ──────────────────────────────────────
    src = src_image or pick_sweep_image(val_images_dir)
    for conf in conf_list:
        conf = str(conf)
        name = f"{exp}_conf{conf.replace('.', '')}"
        dt = runs_root / "detect" / name / "detect.txt"
        imgs = sorted((runs_root / "detect" / name).glob("*.jpg")) if (runs_root / "detect" / name).exists() else []
        mt = runs_root / "detect" / name / "meta.json"
        if dt.exists() and imgs and _meta_ok(mt, "detect-conf", conf, wk, weights):
            t = dt.read_text(encoding="utf-8", errors="ignore")
            print(f"conf={conf} (沿用既有)")
        else:
            t = run(["python", "detect.py", "--weights", weights, "--source", str(src),
                     "--img", str(img), "--conf", conf,
                     "--project", str(runs_root / "detect"), "--name", name, "--exist-ok"])
            (runs_root / "detect" / name).mkdir(parents=True, exist_ok=True)
            dt.write_text(t, encoding="utf-8", errors="ignore")
            _meta_write(mt, "detect-conf", conf, wk, weights)
        m = re.search(r"\d+x\d+\s+(.+),\s+[\d.]+ms", t)
        det = m.group(1).strip() if m else ""
        rows.append([exp, "detect-conf", float(conf), None, None, None, None, det,
                     f"圖欄 runs/detect/{name}"])
        print(f"conf={conf} 檢出: {det}")

    # ── 備份掃參輸出到雲端────────────────────────
    import shutil
    assert backup_dir and Path(backup_dir).is_absolute(), "BACKUP_DIR 非法"
    assert Path(drive_marker).exists(), "Drive 未掛載，無法備份掃參"
    sweep_backup = Path(backup_dir) / "掃參"
    sweep_backup.mkdir(parents=True, exist_ok=True)
    for pattern, src_root in [("_iou*", "val"), ("_conf*", "detect")]:
        for src in (runs_root / src_root).glob(f"{exp}{pattern}"):
            dst = sweep_backup / src.name
            if dst.exists():
                shutil.rmtree(dst)
            shutil.copytree(src, dst)
    print("備份完成，掃參資料夾內容：")
    for p in sorted(sweep_backup.iterdir()):
        print(" ", p.name)

    # ── 寫入掃參紀錄 Excel ────────────────────────────────────
    if Path(drive_marker).exists():
        target = Path(cloud_xlsx)
    else:  # pragma: no cover - 與上一斷言同條件，保留原分支語義
        target = Path(local_xlsx)
        print("Drive 未掛載，掃參紀錄寫入本地（之後需手動合併，避免與雲端表分叉）")
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(target) if target.exists() else openpyxl.Workbook()
    if "掃參紀錄" not in wb.sheetnames:
        ws = wb.create_sheet("掃參紀錄")
        for c, g in enumerate(SWP_G, 1):
            ws.cell(1, c).value = g
        _merge_runs(ws, SWP_G)
        ws.append(SWP_H)
    else:
        ws = wb["掃參紀錄"]
        _cur = [ws.cell(2, c).value for c in range(1, len(SWP_H) + 1)] if ws.max_row >= 2 else []
        if _cur != list(SWP_H):
            for c, g in enumerate(SWP_G, 1):
                ws.cell(1, c).value = g
            _merge_runs(ws, SWP_G)
            _has = ws.max_row > HEADER_ROW and any(
                ws.cell(r, 1).value is not None for r in range(HEADER_ROW + 1, ws.max_row + 1))
            if _has:
                raise ValueError("「掃參紀錄」表頭與程式預期不符（可能被手動改過），為避免錯位寫入請先備份並對齊表頭")
            for c, h in enumerate(SWP_H, 1):
                ws.cell(2, c).value = h
    have = {
        (str(ws.cell(r, 1).value), str(ws.cell(r, 2).value), str(ws.cell(r, 3).value))
        for r in range(HEADER_ROW + 1, ws.max_row + 1)
        if ws.cell(r, 1).value is not None
    }
    n = 0
    for r in rows:
        if (str(r[0]), str(r[1]), str(r[2])) not in have:
            ws.append(r)
            n += 1
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"].cell(1, 1).value is None:
        del wb["Sheet"]
    wb.save(target)
    print(f"掃參紀錄新增 {n} 列 → {target}")
    return rows